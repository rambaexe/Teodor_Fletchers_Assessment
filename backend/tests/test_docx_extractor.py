import zipfile
from pathlib import Path

import docx
import pytest

from app.extractor.docx import DocxExtractor


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    # built on the fly: no binary fixtures in the repo
    d = docx.Document()
    d.core_properties.title = "Quarterly Report"
    d.core_properties.author = "Jane Doe"
    d.add_heading("Summary", level=1)
    d.add_paragraph("Revenue grew this quarter.")
    d.add_paragraph("")  # empty: should be skipped
    d.add_paragraph("First point", style="List Bullet")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Revenue"
    table.cell(1, 0).text, table.cell(1, 1).text = "EU", "1.2m"
    d.add_paragraph("Closing note.")

    path = tmp_path / "report.docx"
    d.save(str(path))
    return path


def header(path: Path) -> bytes:
    return path.read_bytes()[:1024]


def test_can_handle_docx(sample_docx):
    assert DocxExtractor.can_handle(header(sample_docx), sample_docx)


def test_rejects_plain_zip(tmp_path):
    path = tmp_path / "archive.docx"  # misleading extension
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("readme.txt", "hi")
    assert not DocxExtractor.can_handle(header(path), path)


def test_rejects_non_zip(tmp_path):
    path = tmp_path / "fake.docx"
    path.write_bytes(b"%PDF-1.4 not a docx")
    assert not DocxExtractor.can_handle(header(path), path)


def test_extracts_blocks_in_order(sample_docx):
    result = DocxExtractor().extract(sample_docx, on_progress=lambda _: None)

    assert [(c.kind, c.text) for c in result.chunks] == [
        ("heading", "Summary"),
        ("paragraph", "Revenue grew this quarter."),
        ("paragraph", "- First point"),
        ("table", "| Region | Revenue |\n| EU | 1.2m |"),
        ("paragraph", "Closing note."),
    ]
    assert [c.position for c in result.chunks] == [0, 1, 2, 3, 4]
    assert all(c.source == "text" and c.page is None for c in result.chunks)


def test_result_fields(sample_docx):
    result = DocxExtractor().extract(sample_docx, on_progress=lambda _: None)
    assert result.method == "text"
    assert result.page_count is None
    assert result.metadata["title"] == "Quarterly Report"
    assert result.metadata["author"] == "Jane Doe"


def test_reports_progress_to_completion(sample_docx):
    progress: list[float] = []
    DocxExtractor().extract(sample_docx, on_progress=progress.append)
    assert progress == sorted(progress)
    assert progress[-1] == 1.0
