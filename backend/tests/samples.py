"""Sample files built on the fly: no binary fixtures in the repo."""

import io
from pathlib import Path

import docx
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

DIGITAL_TEXT = "This page has a real text layer with plenty of characters."


def pdf_bytes(pages: list[str], title: str | None = None) -> bytes:
    """pages: "text" = digital page, "image" = scanned page (image only), "blank"."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    if title:
        c.setTitle(title)
        c.setAuthor("Jane Doe")
    for kind in pages:
        if kind == "text":
            c.drawString(72, 720, DIGITAL_TEXT)
        elif kind == "image":
            c.drawImage(ImageReader(Image.new("RGB", (200, 100), "white")), 72, 600)
        c.showPage()
    c.save()
    return buf.getvalue()


def make_pdf(path: Path, pages: list[str], title: str | None = None) -> Path:
    path.write_bytes(pdf_bytes(pages, title))
    return path


def docx_bytes() -> bytes:
    d = docx.Document()
    d.core_properties.title = "Quarterly Report"
    d.core_properties.author = "Jane Doe"
    d.add_heading("Summary", level=1)
    d.add_paragraph("Revenue grew this quarter.")
    d.add_paragraph("")  # empty: should be skipped
    d.add_paragraph("First point", style="List Bullet")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Revenue"
    table.cell(1, 0).text, table.cell(1, 1).text = "EU", "1.2m"
    d.add_paragraph("Closing note.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
